import json
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from app.services.demand_service import serialize_demand
from app.standards.rules import get_value
from app.standards.schema import Standard


def as_standard(value):
    return value if isinstance(value, Standard) or value is None else Standard.model_validate(value)


def as_demand(value):
    return value if isinstance(value, dict) else serialize_demand(value)


def display(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def cell_escape(value):
    return display(value).replace("|", "\\|").replace("\n", "<br>")


ZONE_TITLES = {"A": "A 区 · 公共要素", "B": "B 区 · 业务需求要素（Why）",
               "C": "C 区 · 用户需求要素（Who/What）", "D": "D 区 · 功能需求要素（How）"}
SOURCE_LABELS = {"user": "用户", "agent": "AI 提炼", "default": "默认"}


def user_supplements(content: str | None, messages) -> list:
    """提报人补充原话：多轮对话中除 A8 底稿（与 content 全等的首条 USER 消息）外的全部 USER 消息，保持时间序。
    展示层聚合并用（详情页 A8 小节 / 导出 md），A8 底稿语义不动。"""
    base = (content or "").strip()
    skipped, supplements = False, []
    for msg in messages:
        if msg.get("role") != "USER":
            continue
        text = (msg.get("content") or "").strip()
        if not skipped and base and text == base:
            skipped = True
            continue
        supplements.append(msg)
    return supplements


def to_markdown(demand, standard, messages) -> str:
    d, std = as_demand(demand), as_standard(standard)
    lines = [f"# {d.get('title') or '未命名草稿'}", ""]
    # A8 原始提报文本置首（PRD §4.2.1：原话完整保留，作为追溯底稿）
    lines.extend(["## 原始提报文本（A8）", "", d.get("content") or "", ""])
    supplements = user_supplements(d.get("content"), messages)
    if supplements:
        lines.extend(["### 提报人补充原话", ""])
        for msg in supplements:
            text = (msg.get("content") or "").replace("\n", " ")
            lines.append(f"- {msg.get('createdAt', '')}：{text}")
        lines.append("")
    lines.extend(["## 基本信息", "", "| 字段 | 内容 |", "|---|---|"])
    labels = {"id": "ID", "demandNo": "编号", "demandTypeCode": "类型", "subtypeCode": "子类", "status": "状态",
              "urgency": "紧急度", "expectDeliveryAt": "期望交付日期", "submitterName": "提报人",
              "submitterDept": "提报部门", "channel": "渠道", "revision": "修改次数", "createdAt": "创建时间",
              "submittedAt": "提交时间", "closedAt": "撤销时间", "closeReason": "撤销原因",
              "standardVersionId": "标准快照ID", "qualityContentHash": "内容哈希"}
    for key, label in labels.items():
        lines.append(f"| {label} | {cell_escape(d.get(key))} |")
    lines.append("")
    if std and d.get("elements"):
        # TECH 五区表单：按 A/B/C/D 分区排版（PRD §4），顶层公共字段已在基本信息
        for zone in ("A", "B", "C", "D"):
            fields = [f for f in std.fields
                      if f.zone == zone and "." in f.field_path and get_value(d, f.field_path) is not None]
            if not fields:
                continue
            lines.extend([f"## {ZONE_TITLES[zone]}", "", "| 要素 | 内容 |", "|---|---|"])
            for f in fields:
                lines.append(f"| {cell_escape(f.label)} | {cell_escape(get_value(d, f.field_path))} |")
            lines.append("")
    else:
        # MATL/TRAIN ext 表单（未升级五区）
        lines.extend(["## 标准字段", "", "| 字段 | 内容 |", "|---|---|"])
        fields = {f.field_path: f.label for f in std.fields} if std else {}
        for key, value in (d.get("ext") or {}).items():
            lines.append(f"| {cell_escape(fields.get('ext.' + key, key))} | {cell_escape(value)} |")
        lines.append("")
    # E 区：系统维护状态附注（elementStatus/fieldSources/revision/changeLog）
    lines.extend(["## E 区 · 状态信息", "", f"- 需求状态：{d.get('status') or ''}", f"- 修改次数：{d.get('revision') or 0}", ""])
    sources = d.get("fieldSources") or {}
    if sources:
        lines.extend(["### 字段来源", "", "| 字段 | 来源 |", "|---|---|"])
        for path, source in sources.items():
            lines.append(f"| {cell_escape(path)} | {SOURCE_LABELS.get(source, source)} |")
        lines.append("")
    logs = d.get("changeLogs") or []
    if logs:
        lines.extend(["### 变更留痕", "", "| 时间 | 字段 | 旧值 | 新值 | 来源 |", "|---|---|---|---|---|"])
        for log in logs:
            source = SOURCE_LABELS.get(log.get("source"), log.get("source") or "")
            lines.append(f"| {cell_escape(log.get('createdAt'))} | {cell_escape(log.get('fieldKey'))} | {cell_escape(log.get('oldValue'))} | {cell_escape(log.get('newValue'))} | {source} |")
        lines.append("")
    lines.extend(["## 要素完备度", "", "| 要素 | 状态 | 说明 |", "|---|---|---|"])
    element_labels = {f.key: f.label for f in std.elements} if std else {}
    for item in d.get("quality") or []:
        lines.append(f"| {cell_escape(element_labels.get(item['key'], item['key']))} | {cell_escape(item['status'])} | {cell_escape(item.get('note'))} |")
    lines.extend(["", "## 对话回放", ""])
    for msg in messages:
        lines.extend([f"### {msg['role']} · {msg.get('createdAt', '')}", "", msg["content"], ""])
    return "\n".join(lines) + "\n"


def to_xlsx(demands, standards) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    groups = {}
    for row in demands:
        d = as_demand(row)
        groups.setdefault(d.get("demandTypeCode") or "未分类", []).append(d)
    if not groups:
        groups["需求"] = []
    basic = [("id", "ID"), ("demandNo", "编号"), ("title", "标题"), ("demandTypeCode", "类型"), ("subtypeCode", "子类"),
             ("content", "需求描述"), ("status", "状态"), ("revision", "修改次数"), ("urgency", "紧急度"), ("expectDeliveryAt", "期望交付日期"),
             ("submitterName", "提报人"), ("submitterDept", "提报部门"), ("channel", "渠道"),
             ("submittedAt", "提交时间"), ("closedAt", "撤销时间"), ("closeReason", "撤销原因"),
             ("standardVersionId", "标准快照ID"), ("qualityContentHash", "内容哈希")]
    for type_code, rows in groups.items():
        sheet = workbook.create_sheet(type_code)
        fields, elements = {}, {}
        for row in rows:
            std = as_standard(standards.get(row["id"]) or standards.get(type_code))
            if std:
                for f in std.fields:
                    # 非顶层路径（ext.* / elements.区.*）都进入动态列
                    if "." in f.field_path:
                        fields.setdefault(f.field_path, f.label)
                for f in std.elements:
                    elements.setdefault(f.key, f.label)
        sheet.append([label for _, label in basic] + [label + "状态" for label in elements.values()] + list(fields.values()))
        for row in rows:
            quality = {q["key"]: q["status"] for q in row.get("quality") or []}
            values = [row.get(key) for key, _ in basic] + [quality.get(key) for key in elements] + [get_value(row, path) for path in fields]
            sheet.append([display(value) for value in values])
            # 用户文本强制字符串类型，=、+、-、@ 等不会变为 Excel 公式。
            for cell in sheet[sheet.max_row]:
                cell.data_type = "s"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 22
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
