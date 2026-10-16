"""P4 报表导出单测（2026-10-14）。

覆盖口径（对应迁移方案「异步 Excel 导出：任务表 + 后台生成 + /report/download/{taskId}」）：
- task_no RPT-YYYYMMDD-NNN 当日递增；类型仅 DEMAND_LIST
- 状态机 PENDING→RUNNING→SUCCESS/FAILED；xlsx 13 列固定表头、状态中文、交付周期小时
- 查询条件 status/type/urgency/keyword/mine/submittedFrom-To；EXPORT_LIMIT 截断
- 成功后 REPORT_READY 站内信（模板缺失只记日志不影响成功）；失败写 error_msg
- API：export 角色守卫、tasks 仅本人、download 本人或 ADMIN/EXECUTIVE 且 SUCCESS
"""
import asyncio
import json
from datetime import timedelta
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.config import get_settings
from app.core.errors import BizError
from app.db.models import Demand, Notification, NotificationTemplate, ReportExportTask, User, now
from app.services import report_service

# ---- 基座工具 ----

async def make_user(factory, name, wecom_userid, role="HANDLER", type_scope=None):
    # 业务角色默认配全类型权限（贴合生产「必配」约束），按需用 type_scope 覆盖
    if type_scope is None and role in ("EXECUTIVE", "MANAGER", "HANDLER"):
        type_scope = "TECH,MATL,TRAIN"
    async with factory() as db:
        async with db.begin():
            user = User(name=name, wecom_userid=wecom_userid, channel="CHUANGJIN_LS", role=role,
                        type_scope=type_scope)
            db.add(user)
            await db.flush()
            db.expunge(user)
            return user


@pytest.fixture
async def submitter(session_factory):
    return await make_user(session_factory, "提报人", "p4r-submitter")


@pytest.fixture(autouse=True)
async def report_seed(session_factory, tmp_path, monkeypatch):
    """铺 REPORT_READY 模板（conftest 清库会抹掉迁移种子）；报表目录隔离到 tmp_path。"""
    monkeypatch.setattr(get_settings(), "report_dir", str(tmp_path))
    async with session_factory() as db:
        async with db.begin():
            db.add(NotificationTemplate(
                template_code="REPORT_READY", template_name="报表就绪",
                title_template="报表已生成：${report_name}",
                content_template="您导出的${report_name}（${task_no}）已生成，可前往下载。"))
    return tmp_path


async def insert_demand(factory, submitter, minutes_ago=0, **overrides):
    async with factory() as db:
        async with db.begin():
            moment = now() - timedelta(minutes=minutes_ago)
            demand = Demand(
                title="报表需求", demand_type_code="TECH", content="内容", urgency="普通",
                status="SUBMITTED", submitter_id=submitter.id, submitter_name=submitter.name,
                submitter_dept="科技部", channel="CHUANGJIN_LS",
                demand_no=f"D{uuid4().hex[:10]}",
                client_request_id=f"raw-{uuid4().hex}", create_request_hash="raw",
                submitted_at=moment, revision=0, created_at=moment, updated_at=moment)
            for key, value in overrides.items():
                setattr(demand, key, value)
            db.add(demand)
            await db.flush()
            return demand.id


async def create(factory, user, params=None, report_type="DEMAND_LIST"):
    async with factory() as db:
        async with db.begin():
            task = await report_service.create_task(db, user, report_type, params or {})
            return task.id


async def get_task(factory, task_id):
    async with factory() as db:
        task = await db.get(ReportExportTask, task_id)
        return report_service.serialize_task(task)


async def run_export(factory, user, params=None):
    task_id = await create(factory, user, params)
    await report_service.execute(task_id)
    return task_id, await get_task(factory, task_id)


# ---- 1. 任务创建 ----

@pytest.mark.asyncio
async def test_task_no_sequence_and_type_check(session_factory, submitter):
    first = await create(session_factory, submitter)
    second = await create(session_factory, submitter)
    prefix = "RPT-" + now().strftime("%Y%m%d")
    assert (await get_task(session_factory, first))["taskNo"] == f"{prefix}-001"
    assert (await get_task(session_factory, second))["taskNo"] == f"{prefix}-002"
    with pytest.raises(BizError):
        await create(session_factory, submitter, report_type="MONTHLY_REVIEW")


# ---- 2. 后台执行 ----

@pytest.mark.asyncio
async def test_execute_success_xlsx_and_notice(session_factory, submitter, report_seed):
    handler = await make_user(session_factory, "处理人", "p4r-hdl")
    await insert_demand(session_factory, submitter, minutes_ago=120, title="导出的需求",
                        status="DONE", closed_at=now(), assignee_id=handler.id,
                        expect_delivery_at=now().date())
    task_id, task = await run_export(session_factory, submitter)
    assert task["status"] == "SUCCESS" and task["finishedAt"] is not None
    assert task["fileName"] == f"需求清单-{task['taskNo']}.xlsx"
    # xlsx 读回校验：13 列固定表头 + 数据行（状态中文、周期 2h、处理人名）
    workbook = load_workbook(report_seed / f"{task['taskNo']}.xlsx")
    rows = list(workbook.active.iter_rows(values_only=True))
    assert list(rows[0]) == report_service.HEADERS
    assert len(rows) == 2
    row = rows[1]
    assert row[1] == "导出的需求" and row[3] == "已完成" and row[4] == "普通"
    assert row[5] == "提报人" and row[6] == "科技部" and row[8] == "处理人"
    assert row[12] == 2.0  # 交付周期小时
    # REPORT_READY 站内信（写入即送达）
    async with session_factory() as db:
        mail = (await db.execute(select(Notification).where(
            Notification.receiver_id == submitter.id,
            Notification.template_code == "REPORT_READY"))).scalar_one()
    assert mail.send_status == "SENT" and mail.link == "/report"
    assert task["taskNo"] in mail.content


@pytest.mark.asyncio
async def test_query_filters(session_factory, submitter):
    other = await make_user(session_factory, "别人", "p4r-other")
    await insert_demand(session_factory, submitter, title="命中关键字", urgency="紧急")
    await insert_demand(session_factory, submitter, title="不命中", demand_type_code="MATL")
    await insert_demand(session_factory, other, title="命中关键字")  # mine 过滤掉
    await insert_demand(session_factory, submitter, minutes_ago=60 * 24 * 40,
                        title="命中关键字")  # 时间范围过滤掉
    params = {"keyword": "命中", "mine": True, "urgency": "紧急",
              "submittedFrom": (now().date() - timedelta(days=7)).isoformat(),
              "submittedTo": (now().date() + timedelta(days=1)).isoformat(),
              "status": "SUBMITTED", "demandTypeCode": "TECH"}
    _, task = await run_export(session_factory, submitter, params)
    assert task["status"] == "SUCCESS"
    saved = json.loads((await _raw_task(session_factory, task_id=task["id"])).params_json)
    assert saved["keyword"] == "命中"


async def _raw_task(factory, task_id):
    async with factory() as db:
        return await db.get(ReportExportTask, task_id)


@pytest.mark.asyncio
async def test_query_filters_row_count(session_factory, submitter, report_seed):
    other = await make_user(session_factory, "别人", "p4r-o2")
    await insert_demand(session_factory, submitter, title="命中关键字", urgency="紧急")
    await insert_demand(session_factory, submitter, title="不命中")
    await insert_demand(session_factory, other, title="命中关键字", urgency="紧急")
    params = {"keyword": "命中", "mine": True, "urgency": "紧急"}
    task_id, task = await run_export(session_factory, submitter, params)
    workbook = load_workbook(report_seed / f"{task['taskNo']}.xlsx")
    assert workbook.active.max_row == 2  # 表头 + 仅本人命中行


@pytest.mark.asyncio
async def test_export_limit(session_factory, submitter, report_seed, monkeypatch):
    monkeypatch.setattr(report_service, "EXPORT_LIMIT", 2)
    for _ in range(3):
        await insert_demand(session_factory, submitter)
    _, task = await run_export(session_factory, submitter)
    workbook = load_workbook(report_seed / f"{task['taskNo']}.xlsx")
    assert workbook.active.max_row == 3  # 表头 + 上限 2 行


@pytest.mark.asyncio
async def test_execute_failure_marks_failed(session_factory, submitter, monkeypatch):
    def boom(task, rows):
        raise RuntimeError("磁盘写失败")

    monkeypatch.setattr(report_service, "_write_xlsx", boom)
    await insert_demand(session_factory, submitter)
    task_id, task = await run_export(session_factory, submitter)
    assert task["status"] == "FAILED" and "磁盘写失败" in task["errorMsg"]


@pytest.mark.asyncio
async def test_template_missing_still_success(session_factory, submitter):
    async with session_factory() as db:
        async with db.begin():
            template = (await db.execute(select(NotificationTemplate).where(
                NotificationTemplate.template_code == "REPORT_READY"))).scalar_one()
            template.status = "INACTIVE"
    await insert_demand(session_factory, submitter)
    _, task = await run_export(session_factory, submitter)
    assert task["status"] == "SUCCESS"  # 模板停用只记日志，不影响任务
    async with session_factory() as db:
        count = (await db.execute(select(Notification).where(
            Notification.receiver_id == submitter.id))).scalars().all()
    assert count == []


# ---- 3. API ----

async def _login(client, name, wecom_userid):
    response = await client.post("/demandhub-api/auth/dev-login",
                                 json={"name": name, "wecomUserid": wecom_userid})
    assert response.status_code == 200
    return {"Authorization": "Bearer " + response.json()["data"]["accessToken"]}


@pytest.mark.asyncio
async def test_api_export_guard_and_background(client, session_factory, submitter):
    await make_user(session_factory, "经理", "p4r-mgr", role="MANAGER")
    manager = await _login(client, "经理", "p4r-mgr")
    handler = await _login(client, "提报人", "p4r-submitter")
    path = "/demandhub-api/report/export"
    assert (await client.post(path, headers=handler, json={})).json()["code"] == 403
    created = (await client.post(path, headers=manager, json={})).json()
    assert created["code"] == 0 and created["data"]["taskNo"].startswith("RPT-")
    # 后台即发即忘任务：轮询到 SUCCESS（同一事件循环内执行）
    for _ in range(100):
        await asyncio.sleep(0.02)
        tasks = (await client.get("/demandhub-api/report/tasks", headers=manager)).json()["data"]
        if tasks["records"] and tasks["records"][0]["status"] == "SUCCESS":
            break
    assert tasks["records"][0]["status"] == "SUCCESS"
    # tasks 仅本人可见
    other_tasks = (await client.get("/demandhub-api/report/tasks", headers=handler)).json()["data"]
    assert other_tasks["total"] == 0


@pytest.mark.asyncio
async def test_api_download_auth(client, session_factory, submitter):
    await make_user(session_factory, "管理者", "p4r-exec", role="EXECUTIVE")
    await make_user(session_factory, "路人", "p4r-passer")
    await insert_demand(session_factory, submitter)
    task_id, _ = await run_export(session_factory, submitter)  # 手动执行，避免后台竞态
    pending_id = await create(session_factory, submitter)
    owner = await _login(client, "提报人", "p4r-submitter")
    executive_h = await _login(client, "管理者", "p4r-exec")
    other_h = await _login(client, "路人", "p4r-passer")
    base = "/demandhub-api/report/download"
    resp = await client.get(f"{base}/{task_id}", headers=owner)
    assert resp.status_code == 200 and len(resp.content) > 0
    assert "spreadsheetml" in resp.headers["content-type"]
    assert (await client.get(f"{base}/{task_id}", headers=executive_h)).status_code == 200
    assert (await client.get(f"{base}/{task_id}", headers=other_h)).json()["code"] == 403
    assert (await client.get(f"{base}/{pending_id}", headers=owner)).json()["code"] == 409
    assert (await client.get(f"{base}/999999", headers=owner)).json()["code"] == 404
